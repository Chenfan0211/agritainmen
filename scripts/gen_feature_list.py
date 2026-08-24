# -*- coding: utf-8 -*-
"""Generate the customer-facing Zhongxuan Technology feature list."""

from __future__ import annotations

from datetime import date
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "中选科技-功能清单.docx"

GREEN = RGBColor(0x1D, 0x6B, 0x44)
DARK_GREEN = RGBColor(0x13, 0x3F, 0x2B)
CHARCOAL = RGBColor(0x25, 0x2A, 0x27)
GRAY = RGBColor(0x67, 0x70, 0x69)
GOLD = RGBColor(0xB4, 0x87, 0x35)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_FILL = "1D6B44"
SECTION_FILL = "E8F1EB"
ALT_FILL = "F6F8F6"
GOLD_FILL = "F5EFE2"


Feature = tuple[str, str, str, str]


SECTIONS: list[tuple[str, str, list[Feature]]] = [
    (
        "供应链管理后台",
        "面向平台运营与管理人员，统一管理供应商、商品、价格、订单、售后、门店、推广与结算。",
        [
            ("经营分析", "核心经营指标", "汇总交易额、订单量、门店数、商品数、供应商数、用户数等核心经营指标。", "平台运营、管理人员"),
            ("经营分析", "交易趋势分析", "按日、周、月查看交易额与订单量变化趋势，支持经营走势对比。", "平台运营、管理人员"),
            ("经营分析", "品类销售分析", "统计各商品品类的销售额、订单量、销量及占比。", "平台运营、商品运营"),
            ("经营分析", "热销商品排行", "按销量或销售额展示热销商品排行，辅助选品与补货决策。", "平台运营、商品运营"),
            ("经营分析", "实时待办中心", "集中展示待审核供应商、待审核商品、待发货订单及待处理售后。", "平台运营、客服"),
            ("经营分析", "报表查询与导出", "按时间、门店、供应商、商品等条件查询并导出经营报表。", "平台运营、财务人员"),
            ("供应商管理", "供应商邀请入驻", "录入供应商名称、经营品类、地区及联系人信息，发起合作邀请。", "平台运营"),
            ("供应商管理", "供应商资料维护", "维护企业信息、联系人、联系电话、经营区域及供货品类。", "平台运营、供应商"),
            ("供应商管理", "供应商资质审核", "审核营业执照、经营许可、资质有效期及相关证明材料。", "平台运营、审核人员"),
            ("供应商管理", "合作状态管理", "管理待审核、合作中、已停用等供应商合作状态。", "平台运营"),
            ("供应商管理", "供应品类管理", "配置供应商可供商品分类、品牌及服务区域。", "平台运营、商品运营"),
            ("供应商管理", "供应商结算", "统计供应商应结货款、退款扣减、结算状态及历史结算记录。", "平台运营、财务人员"),
            ("商品管理", "门店订货商品", "维护面向农家乐门店采购的供货商品目录。", "商品运营"),
            ("商品管理", "用户商城商品", "维护面向消费者销售的快递直发商品目录。", "商品运营"),
            ("商品管理", "商品基础信息", "维护商品名称、分类、供应商、标签、规格、详情及配送方式。", "商品运营、供应商"),
            ("商品管理", "商品图片与素材", "维护商品主图、轮播图、规格图及商品介绍素材。", "商品运营、供应商"),
            ("商品管理", "SKU 规格管理", "支持一个商品配置多个规格、规格编码、规格价格与独立库存。", "商品运营、供应商"),
            ("商品管理", "库存管理", "查看与调整商品及 SKU 库存，设置库存预警并记录库存变化。", "商品运营、供应商"),
            ("商品管理", "商品审核", "对供应商提交或门店自营商品进行审核、驳回和重新提交。", "平台运营、审核人员"),
            ("商品管理", "商品上下架", "控制商品销售状态及各业务端的展示范围。", "商品运营"),
            ("价格体系", "集采价格", "维护平台集中采购价格与建议零售价。", "商品运营、财务人员"),
            ("价格体系", "阶梯价格", "根据采购数量设置分档价格，采购量达到条件后自动匹配。", "商品运营"),
            ("价格体系", "区域价格", "按城市、区域或服务范围配置差异化销售价格。", "商品运营"),
            ("价格体系", "会员价格", "按会员等级配置专属价格和优惠规则。", "平台运营、商品运营"),
            ("价格体系", "分销价格", "配置商品基础价、一级分佣、二级分佣及用户端展示价格。", "平台运营、财务人员"),
            ("订单履约", "订单统一查询", "统一查询门店采购订单、用户商城订单及各供应商子订单。", "平台运营、客服"),
            ("订单履约", "多供应商拆单", "根据商品所属供应商将主订单拆分为独立履约子订单。", "平台运营、供应商"),
            ("订单履约", "订单状态管理", "管理待支付、待接单、待发货、配送中、已收货、已取消等订单状态。", "平台运营、客服"),
            ("订单履约", "批量发货", "选择多笔符合条件的订单统一执行发货处理。", "平台运营、供应商"),
            ("订单履约", "物流跟踪", "查看配送方式、承运信息、运单号、物流轨迹及签收结果。", "平台运营、客服"),
            ("售后服务", "售后工单", "统一受理退款、退货、补发及质量理赔等售后申请。", "客服、平台运营"),
            ("售后服务", "售后审核", "审核售后原因、商品数量、申请金额及相关凭证。", "客服、审核人员"),
            ("售后服务", "退款与退货", "处理仅退款、退货退款、退款失败重试及退货物流。", "客服、财务人员"),
            ("售后服务", "售后统计", "统计售后数量、退款金额、售后率及供应商责任情况。", "平台运营、客服"),
            ("门店管理", "门店入驻管理", "维护农家乐门店名称、地区、地址、联系人及营业信息。", "平台运营"),
            ("门店管理", "门店经营资料", "维护门店标签、评分、人均消费、展示图片及经营状态。", "平台运营、门店店长"),
            ("门店管理", "门店账号管理", "创建店长与店员账号，设置启停状态及所属门店。", "平台运营、门店店长"),
            ("门店管理", "门店推广权限", "控制店员推广资格及门店推广业务权限。", "平台运营、门店店长"),
            ("推广与结算", "推客资源管理", "维护推客、达人和主播资料、等级、类型及合作状态。", "平台运营"),
            ("推广与结算", "推广效果统计", "统计分享次数、访问用户、绑定用户、成交订单及推广贡献。", "平台运营、推客"),
            ("推广与结算", "分佣规则配置", "按门店、商品、直播及推广角色配置佣金比例和生效状态。", "平台运营、财务人员"),
            ("推广与结算", "佣金结算", "汇总待结算、可结算、已结算佣金并生成结算记录。", "平台运营、财务人员"),
            ("基础数据", "品类管理", "维护商品品类和供应商品类，支持新增、修改、排序与停用。", "平台运营、商品运营"),
            ("基础数据", "业务字典", "维护城市、售后原因、业务状态及其他统一字典数据。", "平台运营"),
            ("基础数据", "通知与操作记录", "管理平台通知、已读状态、关键业务操作记录及数据导出记录。", "平台运营、管理人员"),
        ],
    ),
    (
        "农家乐门店端",
        "面向顾客、店员和店长，承接门店展示、预订、商城、会员、推广及日常经营管理。",
        [
            ("多门店经营", "独立门店配置", "为每家农家乐配置独立名称、品牌标识、主题风格和访问入口。", "平台运营、门店店长"),
            ("多门店经营", "门店联系信息", "展示联系电话、详细地址、营业时间及导航入口。", "顾客、门店人员"),
            ("门店展示", "门店首页", "展示门店形象、推荐内容、营销公告和快捷服务入口。", "顾客"),
            ("门店展示", "环境与相册", "展示门店环境、包厢、菜品、服务场地及活动图片。", "顾客"),
            ("门店展示", "招牌土菜", "展示菜品名称、介绍、图片、现价及原价。", "顾客"),
            ("门店展示", "特色包厢", "展示包厢名称、容纳人数、最低消费、使用状态及介绍。", "顾客"),
            ("门店展示", "特色服务", "展示采摘、柴火灶、露营等服务内容、适合人群、时长及价格。", "顾客"),
            ("预约预订", "包厢预订", "选择包厢、日期、场次、人数并提交预订。", "顾客"),
            ("预约预订", "套餐预订", "查看套餐内容、价格和适用人数后提交预订。", "顾客"),
            ("预约预订", "体验服务下单", "选择特色服务、体验日期和参与人数，自动计算费用并下单。", "顾客"),
            ("预约预订", "重复预订校验", "对相同日期、场次和服务的重复预订进行提醒和限制。", "顾客"),
            ("预约预订", "预订记录", "查看预订内容、时间、人数、金额及当前状态。", "顾客、门店人员"),
            ("预约预订", "取消与提醒", "支持取消符合条件的预订，并发送到店及状态变化提醒。", "顾客、门店人员"),
            ("预约预订", "到店核销", "店员或店长核验预订信息并完成到店核销。", "店员、门店店长"),
            ("门店商城", "商品浏览与搜索", "按分类浏览商品，支持关键词搜索及商品筛选。", "顾客"),
            ("门店商城", "商品详情与规格", "查看商品图文详情、供应来源、价格、库存及规格。", "顾客"),
            ("门店商城", "购物车", "支持加入购物车、规格区分、数量调整、删除和金额汇总。", "顾客"),
            ("门店商城", "库存校验与结算", "结算时校验商品库存、会员价格和订单金额。", "顾客"),
            ("门店商城", "门店订单", "查看订单商品、金额、支付状态、配送状态和售后入口。", "顾客、门店人员"),
            ("会员中心", "会员等级与会员卡", "展示会员编号、会员等级、余额、积分及会员权益。", "顾客"),
            ("会员中心", "余额充值与消费", "支持会员余额充值、余额支付及余额变动记录。", "顾客"),
            ("会员中心", "积分管理", "记录消费积分、活动积分及积分变动明细。", "顾客、门店店长"),
            ("角色权限", "顾客角色", "提供门店浏览、预订、购物、会员和个人订单功能。", "顾客"),
            ("角色权限", "店员角色", "提供预订查询、到店核销及授权后的推广功能。", "店员"),
            ("角色权限", "店长角色", "提供门店资料、包厢、菜品、服务和员工账号管理功能。", "门店店长"),
            ("经营管理", "包厢管理", "新增、编辑和删除包厢，维护人数、价格、图片及使用状态。", "门店店长"),
            ("经营管理", "菜品管理", "新增、编辑和删除招牌菜品，维护价格、介绍及图片。", "门店店长"),
            ("经营管理", "特色服务管理", "维护特色服务内容、价格、适用人群、时长及可预订日期。", "门店店长"),
            ("经营管理", "店员账号管理", "新增、修改、启停店员账号并分配角色和推广权限。", "门店店长"),
            ("推广分成", "店员推广码", "授权店员生成专属推广链接、二维码及分享素材。", "店员、门店店长"),
            ("推广分成", "用户绑定与消费分成", "记录推广来源、绑定用户、归因订单及门店消费分成。", "店员、门店店长"),
            ("账号与个人中心", "登录与退出", "支持微信授权、门店账号登录、安全退出和登录状态管理。", "顾客、店员、门店店长"),
        ],
    ),
    (
        "农家乐联盟端",
        "面向消费者和联盟推客，提供门店发现、榜单线路、直播活动、预订及推广佣金服务。",
        [
            ("附近门店", "定位与城市切换", "获取用户位置并支持手动切换城市，展示对应区域门店。", "用户、推客"),
            ("附近门店", "地图分布", "在地图中展示附近农家乐位置、距离及快捷查看入口。", "用户、推客"),
            ("附近门店", "距离排序", "按用户当前位置与门店距离由近到远展示。", "用户、推客"),
            ("附近门店", "综合筛选", "按好评、人气、直播状态和可预订状态筛选门店。", "用户、推客"),
            ("门店榜单", "多维榜单", "提供人气榜、好评榜和直播榜，展示排名及关键指标。", "用户、推客"),
            ("乡村线路", "线路展示", "展示乡村旅游线路、行程内容、关联门店和线路特色。", "用户、推客"),
            ("乡村线路", "线路报名", "选择线路并提交参加信息，查看已报名线路。", "用户"),
            ("直播活动", "直播列表", "展示正在直播和即将开播的活动、主播及观看人数。", "用户、推客"),
            ("直播活动", "开播提醒", "对直播预告设置提醒，并在开播时发送通知。", "用户、推客"),
            ("直播活动", "观看与分享", "进入直播活动并分享直播链接或小程序卡片。", "用户、推客"),
            ("门店预订", "门店详情", "查看门店介绍、标签、评分、人均消费、套餐及可预订状态。", "用户"),
            ("门店预订", "在线预订", "选择日期、时段和人数提交门店预订。", "用户"),
            ("推广中心", "多场景推广", "生成门店、商品和直播专属推广链接及分享素材。", "推客"),
            ("推广中心", "锁粉与归因", "记录用户首次访问来源、绑定关系、归因订单和推广贡献。", "推客、平台运营"),
            ("推广中心", "推广效果", "统计分享次数、锁定用户、成交订单和预计佣金。", "推客"),
            ("佣金钱包", "佣金账户", "展示累计佣金、待结算佣金、可提现佣金及佣金明细。", "推客"),
            ("佣金钱包", "佣金提现", "选择提现方式并提交提现申请，查看提现处理记录。", "推客"),
            ("账号中心", "登录与个人中心", "支持手机号密码、短信验证码登录及个人资料管理。", "用户、推客"),
        ],
    ),
    (
        "推客运营端",
        "面向签约推客、达人和主播，管理直播内容、推广素材、绑定用户及消费分成。",
        [
            ("推客工作台", "推客资料", "展示推客身份、等级、联系方式及合作状态。", "推客"),
            ("推客工作台", "经营概览", "汇总直播数量、分享次数、绑定用户、成交金额和分成收益。", "推客"),
            ("直播管理", "创建直播", "填写直播标题、封面、活动状态并创建直播活动。", "推客、主播"),
            ("直播管理", "编辑直播", "修改直播标题、封面、关联门店和关联套餐。", "推客、主播"),
            ("直播管理", "直播状态", "管理直播预告、发布、开播、结束和下架状态。", "推客、主播"),
            ("直播管理", "直播下架与删除", "对不再使用的直播进行下架确认和删除处理。", "推客、主播"),
            ("直播选品", "关联多个门店", "单场直播可选择多个农家乐门店作为推广对象。", "推客、主播"),
            ("直播选品", "关联多个套餐", "为每个门店选择一个或多个套餐券作为直播商品。", "推客、主播"),
            ("内容素材", "直播封面", "选择预置封面或上传图片作为直播活动封面。", "推客、主播"),
            ("推广分发", "推广链接与二维码", "生成用户商城直播入口链接、二维码及小程序分享信息。", "推客、主播"),
            ("推广分发", "链接复制与分享", "复制直播链接并通过社交渠道进行分发。", "推客、主播"),
            ("推广统计", "绑定与分成", "查看分享访问、绑定用户、成交订单及消费分成记录。", "推客"),
        ],
    ),
    (
        "门店订货商城",
        "面向农家乐门店采购人员，提供平台供货商品直采、采购订单及物流售后服务。",
        [
            ("商品采购", "供货商品目录", "展示平台供应商品的分类、名称、图片、供货方及可采购状态。", "门店店长、采购人员"),
            ("商品采购", "商品搜索与筛选", "按关键词、商品分类和供应商筛选采购商品。", "门店店长、采购人员"),
            ("商品采购", "商品详情", "查看规格、供货价、建议零售价、库存、销量和商品介绍。", "门店店长、采购人员"),
            ("商品采购", "多规格采购", "选择不同商品规格及采购数量加入进货车。", "门店店长、采购人员"),
            ("进货车", "进货车管理", "支持数量增减、删除、规格区分、库存校验和金额汇总。", "门店店长、采购人员"),
            ("进货车", "采购备注", "提交采购订单时填写包装、配送及其他采购要求。", "门店店长、采购人员"),
            ("采购订单", "提交采购订单", "根据采购商品生成订单并分配至对应供应商。", "门店店长、采购人员"),
            ("采购订单", "订单分类", "按待接单、待发货、已发货、配送中、已收货和已完成筛选订单。", "门店店长、采购人员"),
            ("采购订单", "订单详情", "查看商品、数量、金额、节省金额、备注及履约进度。", "门店店长、采购人员"),
            ("采购订单", "物流轨迹", "查看接单、发货、配送、签收和完成等物流节点。", "门店店长、采购人员"),
            ("采购订单", "确认收货", "对配送完成的采购订单确认收货并记录签收时间。", "门店店长、采购人员"),
            ("采购订单", "取消订单", "在允许取消的状态下提交采购订单取消操作。", "门店店长、采购人员"),
            ("采购订单", "采购售后", "对已收货商品发起退款、退货或质量问题售后。", "门店店长、采购人员"),
            ("采购订单", "再次采购", "将历史订单中的可采购商品快速加入进货车。", "门店店长、采购人员"),
            ("采购分析", "采购概览", "统计月采购额、采购订单数、待收货数量及节省金额。", "门店店长、采购人员"),
            ("账号中心", "登录与门店身份", "支持手机号密码、短信验证码登录并展示当前门店资料。", "门店店长、采购人员"),
        ],
    ),
    (
        "用户商城端",
        "面向消费者及分销用户，提供商品购买、订单履约、直播活动、售后与佣金服务。",
        [
            ("账号与身份", "微信授权登录", "通过微信授权完成用户登录、身份识别和账号建立。", "用户、分销商"),
            ("账号与身份", "用户会话", "保存用户购物车、身份、推广来源及个人业务数据。", "用户、分销商"),
            ("账号与身份", "分销身份", "支持普通用户、一级分销商和二级分销商身份及权益。", "用户、分销商"),
            ("商品商城", "商品分类", "按土特产、粮油、茶饮等分类浏览快递直发商品。", "用户、分销商"),
            ("商品商城", "商品搜索", "按商品名称、供应商和商品标签搜索。", "用户、分销商"),
            ("商品商城", "商品详情", "展示商品图片、供应商、标签、规格、库存和配送说明。", "用户、分销商"),
            ("商品商城", "分销价格", "根据用户身份展示基础价格及对应分销价格。", "用户、分销商"),
            ("购物车", "加入购物车", "选择商品规格及数量加入个人购物车。", "用户、分销商"),
            ("购物车", "购物车管理", "支持数量调整、商品删除、失效提示、库存校验和金额计算。", "用户、分销商"),
            ("收货地址", "地址维护", "新增、编辑和删除收货人、手机号、地区及详细地址。", "用户、分销商"),
            ("收货地址", "默认地址", "设置默认收货地址并在结算时切换使用地址。", "用户、分销商"),
            ("订单交易", "提交订单", "确认商品、地址、配送方式、订单金额和备注后提交主订单。", "用户、分销商"),
            ("订单交易", "多供应商拆单", "根据商品供应商将主订单拆分为多个独立子订单。", "用户、分销商"),
            ("订单交易", "在线支付", "支持微信支付和其他配置的在线支付方式。", "用户、分销商"),
            ("订单交易", "订单状态", "展示待支付、备货、部分发货、已发货、部分收货、已收货及售后状态。", "用户、分销商"),
            ("订单交易", "取消订单", "对符合条件的待支付或未履约订单发起取消。", "用户、分销商"),
            ("履约售后", "子订单物流", "按供应商子订单查看承运信息、运单号和物流轨迹。", "用户、分销商"),
            ("履约售后", "确认收货", "对已送达的供应商子订单确认收货并完成交易。", "用户、分销商"),
            ("履约售后", "按子订单售后", "针对指定供应商子订单发起退款、退货或质量问题售后。", "用户、分销商"),
            ("直播专区", "直播活动入口", "展示直播活动、主播、关联门店及活动状态。", "用户、分销商"),
            ("直播专区", "直播关联商品", "查看直播关联门店、套餐券和推广商品并进入购买流程。", "用户、分销商"),
            ("推广佣金", "分销关系绑定", "记录用户推广来源、上级分销关系及关系锁定结果。", "用户、分销商"),
            ("推广佣金", "佣金明细", "查看订单来源、佣金金额、待结算、可提现及已结算状态。", "分销商"),
            ("推广佣金", "佣金提现", "汇总可提现佣金并提交提现申请。", "分销商"),
        ],
    ),
    (
        "供应商与司机端",
        "面向供应商和配送司机，统一处理门店采购及用户商城订单的接单、发货、配送与交接。",
        [
            ("账号与权限", "双角色登录", "支持供应商账号和司机账号分别登录对应工作台。", "供应商、司机"),
            ("账号与权限", "供应商数据隔离", "供应商只能查看和处理归属本供应商的订单、司机及履约数据。", "供应商"),
            ("供应商工作台", "订单指标", "展示待接单、待发货、配送中、已完成及缺货订单数量。", "供应商"),
            ("供应商工作台", "履约概览", "汇总订单金额、配送任务、司机状态及异常待办。", "供应商"),
            ("订单管理", "多来源订单", "统一承接门店采购订单和用户商城快递订单。", "供应商"),
            ("订单管理", "订单查询", "按订单号、门店、收货人和关键词查询订单。", "供应商"),
            ("订单管理", "状态与缺货筛选", "按履约状态或是否存在缺货筛选订单。", "供应商"),
            ("订单管理", "接单与批量接单", "确认单笔订单或批量确认多笔待接订单。", "供应商"),
            ("发货配送", "快递直发", "选择快递配送方式并录入承运公司和运单号。", "供应商"),
            ("发货配送", "司机配送", "选择可用司机执行门店配送任务。", "供应商"),
            ("发货配送", "指派与改派司机", "在允许的履约节点指派或更换配送司机。", "供应商"),
            ("发货配送", "快递签收", "接收物流签收结果并更新订单履约状态。", "供应商"),
            ("缺货处理", "缺货登记", "按商品规格登记订购数量、实际数量和缺货数量。", "供应商"),
            ("缺货处理", "缺货明细", "查看订单缺货商品、数量、处理状态和相关记录。", "供应商"),
            ("缺货处理", "补发处理", "对缺货商品创建补发安排并更新处理结果。", "供应商"),
            ("交接管理", "出库交接", "记录供应商出库、司机领货、实际数量、缺货情况和交接时间。", "供应商、司机"),
            ("交接管理", "到店交接", "司机到店后登记交接结果、备注和完成时间。", "司机、门店人员"),
            ("交接管理", "交接日志", "按订单和司机查看出库、领货、到店及异常交接记录。", "供应商、司机"),
            ("司机管理", "司机账号", "新增司机并维护姓名、手机号、登录账号和初始密码。", "供应商"),
            ("司机管理", "司机资料与状态", "编辑司机资料、启停账号并查看当前任务状态。", "供应商"),
            ("司机管理", "司机密码重置", "重置司机登录密码并保留相关操作记录。", "供应商"),
            ("司机任务", "今日配送任务", "展示司机当天待领货、配送中及待交接任务。", "司机"),
            ("司机任务", "历史任务", "查看历史配送订单、完成时间及交接结果。", "司机"),
            ("司机任务", "我的交接", "汇总当前司机参与的出库和到店交接记录。", "司机"),
            ("跨端协同", "履约状态回写", "将接单、发货、运单、配送、签收和售后状态同步至相关业务端。", "供应商、平台运营、用户"),
        ],
    ),
    (
        "平台上线基础能力",
        "为中选科技各业务端提供统一、安全、稳定的数据、交易、消息、地图、文件及运维服务。",
        [
            ("业务服务", "统一业务后台", "提供统一业务服务承接各应用端的查询、交易和管理请求。", "平台运营、各业务端"),
            ("业务服务", "统一数据库", "集中存储用户、门店、商品、订单、库存、结算和履约数据。", "平台运营、各业务端"),
            ("业务服务", "接口服务", "为管理后台、小程序和移动端提供统一业务接口。", "各业务端"),
            ("业务服务", "多端实时同步", "保证商品、库存、订单、物流和佣金数据在各业务端及时同步。", "平台运营、各业务端"),
            ("多租户与权限", "门店数据隔离", "按门店隔离经营资料、员工、预订、会员和订单数据。", "平台运营、门店人员"),
            ("多租户与权限", "供应商数据隔离", "按供应商隔离商品、库存、订单、司机和结算数据。", "平台运营、供应商"),
            ("多租户与权限", "用户数据隔离", "按用户隔离购物车、地址、订单、绑定关系和佣金数据。", "用户、分销商"),
            ("多租户与权限", "角色权限控制", "按平台、门店、推客、供应商和司机角色控制菜单、数据和操作权限。", "平台运营、各业务角色"),
            ("认证安全", "微信授权", "接入微信身份授权并建立平台用户账号。", "用户、门店顾客"),
            ("认证安全", "账号密码登录", "支持平台、门店、供应商和司机账号密码登录。", "平台运营、门店人员、供应商、司机"),
            ("认证安全", "短信验证码", "支持登录、身份验证及关键操作所需的短信验证码。", "各业务角色"),
            ("认证安全", "登录态与权限校验", "管理登录有效期、退出、异常失效及接口访问权限。", "各业务角色"),
            ("支付结算", "微信支付", "接入微信支付完成商品、预订和会员充值交易。", "用户、门店顾客"),
            ("支付结算", "余额支付", "支持会员余额抵扣和余额支付。", "门店顾客"),
            ("支付结算", "退款处理", "支持原路退款、部分退款、退款查询和失败重试。", "用户、客服、财务人员"),
            ("支付结算", "分账与结算", "按订单规则计算平台、门店、供应商和推广角色的应结金额。", "平台运营、财务人员"),
            ("支付结算", "佣金提现", "处理推客和分销商提现申请、审核、打款及结果通知。", "推客、分销商、财务人员"),
            ("库存交易", "库存锁定", "用户提交订单后锁定对应商品库存，防止重复占用。", "平台运营、供应商"),
            ("库存交易", "库存扣减与释放", "在支付、取消、超时及售后场景下执行库存扣减或释放。", "平台运营、供应商"),
            ("库存交易", "库存流水", "记录入库、出库、锁定、释放、调整及售后库存变化。", "平台运营、供应商"),
            ("库存交易", "超卖与重复处理控制", "通过交易校验保证订单、支付和库存操作准确执行一次。", "平台运营、供应商"),
            ("物流配送", "快递接口", "对接快递服务，查询承运公司、运单及配送状态。", "平台运营、供应商、用户"),
            ("物流配送", "电子面单", "生成和打印快递电子面单，关联订单与包裹。", "供应商"),
            ("物流配送", "物流订阅与回调", "订阅物流轨迹并接收揽收、运输、派送和签收状态。", "平台运营、供应商、用户"),
            ("消息通知", "短信通知", "发送登录验证码、订单、配送、售后和提现通知。", "各业务角色"),
            ("消息通知", "微信订阅消息", "发送预订提醒、开播提醒、订单状态和售后结果通知。", "用户、门店顾客、推客"),
            ("消息通知", "站内通知", "在各业务端展示待办、业务提醒和系统公告。", "各业务角色"),
            ("地图服务", "地图定位与地址解析", "获取位置并完成地址选择、经纬度解析和地图展示。", "用户、门店人员、司机"),
            ("地图服务", "距离与线路导航", "计算用户、门店和配送地址距离并提供导航线路。", "用户、推客、司机"),
            ("地图服务", "配送服务范围", "配置供应商、门店和司机的配送区域及服务半径。", "平台运营、供应商"),
            ("文件与素材", "业务文件存储", "统一存储商品图片、门店相册、资质文件、售后凭证和直播封面。", "平台运营、各业务角色"),
            ("文件与素材", "二维码与推广素材", "生成推广二维码、直播二维码、门店码及分享素材。", "平台运营、推客、店员"),
            ("审计与报表", "操作与登录日志", "记录用户登录、关键业务操作、权限变更和异常访问。", "平台运营、安全管理人员"),
            ("审计与报表", "审计记录", "保留订单、支付、退款、库存、结算和权限变更审计链路。", "平台运营、财务人员"),
            ("审计与报表", "数据导出与统计", "按权限导出业务数据并生成经营、财务和履约统计报表。", "平台运营、财务人员"),
            ("运行保障", "数据备份与恢复", "定期备份业务数据并支持故障后的数据恢复。", "平台运维"),
            ("运行保障", "运行监控与告警", "监控系统、接口、任务和第三方服务状态并发送异常告警。", "平台运维"),
            ("运行保障", "接口日志与问题追踪", "记录接口请求、异常信息和链路标识，支持问题定位。", "平台运维、技术支持"),
            ("隐私与安全", "隐私授权", "在采集位置、手机号和身份信息前完成用户授权与用途告知。", "用户、各业务角色"),
            ("隐私与安全", "敏感数据保护", "对密码、手机号、身份证明及支付信息进行加密和脱敏。", "平台运营、安全管理人员"),
            ("隐私与安全", "访问与安全防护", "提供访问控制、接口防护、频率限制、风险识别及安全策略。", "平台运维、安全管理人员"),
            ("系统配置", "业务参数", "配置订单时效、库存规则、分佣比例、售后规则及业务开关。", "平台运营"),
            ("系统配置", "第三方参数", "配置支付、物流、短信、地图、文件存储及消息模板参数。", "平台运营、平台运维"),
        ],
    ),
]


def set_font(run, size: float, *, bold: bool = False, color: RGBColor = CHARCOAL, font: str = "微软雅黑") -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 70, start: int = 90, bottom: int = 70, end: int = 90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_text(cell, text: str, *, size: float = 8.7, bold: bool = False, color: RGBColor = CHARCOAL, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.12
    run = paragraph.add_run(text)
    set_font(run, size, bold=bold, color=color)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    current = OxmlElement("w:t")
    current.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, current, end):
        run._r.append(element)
    set_font(run, 8.5, color=GRAY)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_font(run, 15 if level == 1 else 11.5, bold=True, color=GREEN if level == 1 else DARK_GREEN)
    if level == 1:
        p_pr = paragraph._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "10")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "B48735")
        borders.append(bottom)
        p_pr.append(borders)


def add_body(doc: Document, text: str, *, bold: bool = False, color: RGBColor = CHARCOAL, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after: float = 5) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.4
    run = paragraph.add_run(text)
    set_font(run, 10.5, bold=bold, color=color)


def configure_document(doc: Document) -> None:
    doc.core_properties.title = "中选科技功能清单"
    doc.core_properties.subject = "中选科技正式上线功能范围说明"
    doc.core_properties.author = "中选科技"
    doc.core_properties.keywords = "中选科技,功能清单,供应链,农家乐,商城,履约"

    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal.font.size = Pt(10.5)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")

    section = doc.sections[0]
    configure_page(section)
    section.different_first_page_header_footer = True


def configure_page(section) -> None:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.75)


def configure_body_header_footer(section) -> None:
    configure_page(section)

    header = section.header
    header.is_linked_to_previous = False
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header_p.add_run("中选科技  |  功能范围说明")
    set_font(header_run, 8.5, bold=True, color=GREEN)

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run("中选科技功能清单  ·  第 ")
    set_font(footer_run, 8.5, color=GRAY)
    add_page_number(footer_p)
    footer_end = footer_p.add_run(" 页")
    set_font(footer_end, 8.5, color=GRAY)

    sect_pr = section._sectPr
    page_number_type = sect_pr.find(qn("w:pgNumType"))
    if page_number_type is None:
        page_number_type = OxmlElement("w:pgNumType")
        sect_pr.append(page_number_type)
    page_number_type.set(qn("w:start"), "1")


def add_cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()

    eyebrow = doc.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    eyebrow.paragraph_format.space_after = Pt(12)
    run = eyebrow.add_run("ZHONGXUAN TECHNOLOGY")
    set_font(run, 10, bold=True, color=GOLD)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("中选科技")
    set_font(run, 28, bold=True, color=GREEN)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("功能清单")
    set_font(run, 20, bold=True, color=CHARCOAL)

    line_table = doc.add_table(rows=1, cols=1)
    line_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    line_table.autofit = False
    line_table.columns[0].width = Cm(7)
    shade(line_table.cell(0, 0), "B48735")
    line_table.cell(0, 0).height = Cm(0.08)
    line_table.cell(0, 0).text = ""

    tagline = doc.add_paragraph()
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tagline.paragraph_format.space_before = Pt(18)
    tagline.paragraph_format.space_after = Pt(90)
    run = tagline.add_run("农家乐数字供应链、门店经营、联盟引流、分销商城与配送履约一体化平台")
    set_font(run, 11, color=GRAY)

    info = doc.add_table(rows=3, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    info.columns[0].width = Cm(3.5)
    info.columns[1].width = Cm(9.5)
    cover_info = [
        ("项目名称", "中选科技"),
        ("文档名称", "中选科技功能清单"),
        ("编制日期", f"{date.today().year} 年 {date.today().month} 月 {date.today().day} 日"),
    ]
    for row, (label, value) in zip(info.rows, cover_info):
        set_cell_text(row.cells[0], label, size=10, bold=True, color=DARK_GREEN, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[1], value, size=10, align=WD_ALIGN_PARAGRAPH.LEFT)
        shade(row.cells[0], SECTION_FILL)
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_body_header_footer(body_section)


def add_overview(doc: Document) -> None:
    add_heading(doc, "一、项目概述")
    add_body(
        doc,
        "中选科技围绕农家乐经营主体，建设覆盖供应链集采、门店经营、联盟引流、直播推广、用户商城、配送履约、售后结算和数据服务的一体化平台。平台连接运营方、农家乐门店、供应商、推客、消费者和配送司机，形成从商品供给到消费履约的完整业务闭环。",
    )
    add_body(
        doc,
        "本清单面向客户、合作方及项目相关单位，用于说明中选科技正式上线后的完整功能范围。具体功能配置、业务规则和交付边界以最终需求确认结果为准。",
        color=GRAY,
    )

    add_heading(doc, "二、平台组成")
    components = [
        ("01", "供应链管理后台", "统一管理供应商、商品、价格、订单、售后、门店、推广与结算。"),
        ("02", "农家乐门店端", "服务顾客、店员和店长，承接门店经营、预订、商城、会员与推广。"),
        ("03", "农家乐联盟端", "提供附近门店、榜单线路、直播活动、在线预订和推客佣金。"),
        ("04", "推客运营端", "管理直播内容、关联门店套餐、推广链接、绑定用户和消费分成。"),
        ("05", "门店订货商城", "支持农家乐门店按供货价直采、跟踪采购订单并处理售后。"),
        ("06", "用户商城端", "提供分销价格、商品交易、多供应商履约、直播活动与佣金服务。"),
        ("07", "供应商与司机端", "完成接单、发货、缺货、司机配送及出入库交接。"),
        ("08", "平台上线基础能力", "提供数据、权限、支付、库存、物流、消息、安全和运维保障。"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.3, 4.2, 11.5]
    for i, width in enumerate(widths):
        table.columns[i].width = Cm(width)
    for i, title in enumerate(("序号", "平台组成", "主要定位")):
        set_cell_text(table.rows[0].cells[i], title, size=9.2, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(table.rows[0].cells[i], HEADER_FILL)
    set_repeat_table_header(table.rows[0])
    for index, (no, name, desc) in enumerate(components):
        row = table.add_row()
        prevent_row_split(row)
        set_cell_text(row.cells[0], no, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[1], name, bold=True, color=DARK_GREEN)
        set_cell_text(row.cells[2], desc)
        if index % 2:
            for cell in row.cells:
                shade(cell, ALT_FILL)

    add_heading(doc, "三、核心业务闭环")
    stages = [
        ("1", "商品管理", "统一商品、价格与库存"),
        ("2", "推广获客", "门店、联盟、直播与分销"),
        ("3", "用户下单", "预订、采购与商城交易"),
        ("4", "供应商履约", "接单、发货、配送与交接"),
        ("5", "售后结算", "退款、退货、分账与佣金"),
        ("6", "数据分析", "经营、商品、履约与财务分析"),
    ]
    flow = doc.add_table(rows=2, cols=6)
    flow.alignment = WD_TABLE_ALIGNMENT.CENTER
    flow.autofit = False
    for column in flow.columns:
        column.width = Cm(2.75)
    for index, (no, name, desc) in enumerate(stages):
        set_cell_text(flow.cell(0, index), f"{no}  {name}", size=8.7, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(flow.cell(0, index), HEADER_FILL)
        set_cell_text(flow.cell(1, index), desc, size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(flow.cell(1, index), SECTION_FILL if index % 2 == 0 else GOLD_FILL)


def add_feature_table(doc: Document, rows: list[Feature], start_number: int) -> int:
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [0.9, 2.65, 3.05, 7.1, 2.85]
    for index, width in enumerate(widths):
        table.columns[index].width = Cm(width)
    headers = ("序号", "一级模块", "二级功能", "功能说明", "使用对象")
    for index, title in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], title, size=8.8, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(table.rows[0].cells[index], HEADER_FILL)
    set_repeat_table_header(table.rows[0])

    first_cells: dict[str, list] = {}
    for offset, (module, feature, description, audience) in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        values = (str(start_number + offset), module, feature, description, audience)
        for index, value in enumerate(values):
            align = WD_ALIGN_PARAGRAPH.CENTER if index in (0, 1, 4) else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(row.cells[index], value, size=8.35 if index == 3 else 8.2, bold=index == 2, color=DARK_GREEN if index == 2 else CHARCOAL, align=align)
        if offset % 2:
            for cell in row.cells:
                shade(cell, ALT_FILL)
        first_cells.setdefault(module, []).append(row.cells[1])

    for module, module_cells in first_cells.items():
        if len(module_cells) < 2:
            shade(module_cells[0], SECTION_FILL)
            continue
        merged = module_cells[0]
        for cell in module_cells[1:]:
            merged = merged.merge(cell)
        set_cell_text(merged, module, size=8.4, bold=True, color=DARK_GREEN, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(merged, SECTION_FILL)
    return start_number + len(rows)


def add_feature_sections(doc: Document) -> int:
    doc.add_page_break()
    add_heading(doc, "四、详细功能清单")
    add_body(doc, "以下功能按系统端划分，序号全文连续。各模块共同组成中选科技正式上线后的完整产品能力。", color=GRAY)
    next_number = 1
    for section_index, (name, summary, features) in enumerate(SECTIONS, start=1):
        if section_index > 1:
            doc.add_page_break()
        add_heading(doc, f"4.{section_index}  {name}", level=2)
        add_body(doc, summary, color=GRAY, after=7)
        next_number = add_feature_table(doc, features, next_number)
    return next_number - 1


def add_closing(doc: Document, feature_count: int) -> None:
    add_heading(doc, "五、文档说明")
    notes = [
        f"本清单共收录 {feature_count} 项功能，覆盖七个业务应用端及平台上线基础能力。",
        "各系统端通过统一账号、权限、商品、订单、库存、支付、物流、消息和结算服务协同运行。",
        "功能名称与说明用于明确产品能力，具体页面、字段、规则、审批流程及第三方服务接入方式以最终需求确认结果为准。",
        "项目建设过程中新增或调整的功能，应通过正式需求确认流程纳入本清单。",
    ]
    for index, note in enumerate(notes, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.35
        marker = paragraph.add_run(f"{index}. ")
        set_font(marker, 10.2, bold=True, color=GREEN)
        run = paragraph.add_run(note)
        set_font(run, 10.2, color=CHARCOAL)


def validate_source_data() -> int:
    feature_count = sum(len(features) for _, _, features in SECTIONS)
    required_sections = {
        "供应链管理后台",
        "农家乐门店端",
        "农家乐联盟端",
        "推客运营端",
        "门店订货商城",
        "用户商城端",
        "供应商与司机端",
        "平台上线基础能力",
    }
    actual_sections = {name for name, _, _ in SECTIONS}
    if actual_sections != required_sections:
        raise ValueError("Document section set is incomplete")
    if feature_count < 100:
        raise ValueError("Feature list is unexpectedly short")
    for section, _, features in SECTIONS:
        if not features:
            raise ValueError(f"Section has no features: {section}")
        for row in features:
            if len(row) != 4 or any(not value.strip() for value in row):
                raise ValueError(f"Invalid feature row in {section}: {row}")
    return feature_count


def validate_output(path: Path, feature_count: int) -> None:
    rendered = Document(path)
    text_parts = [paragraph.text for paragraph in rendered.paragraphs]
    for table in rendered.tables:
        for row in table.rows:
            text_parts.extend(cell.text for cell in row.cells)
    full_text = "\n".join(text_parts)

    forbidden_terms = (
        "湖南农家乐数字供应链与引流服务平台",
        "演示",
        "Mock",
        "MOCK",
        "未完成",
        "待开发",
        "前端模拟",
        "报价",
        "付款比例",
        "平台技术服务费",
    )
    found = [term for term in forbidden_terms if term in full_text]
    if found:
        raise ValueError(f"Forbidden terms found in output: {found}")
    if "中选科技功能清单" not in full_text:
        raise ValueError("Document title is missing")
    for section, _, _ in SECTIONS:
        if section not in full_text:
            raise ValueError(f"Document section is missing: {section}")

    numbers: list[int] = []
    for table in rendered.tables:
        if not table.rows or len(table.columns) != 5 or table.cell(0, 0).text != "序号":
            continue
        for row in table.rows[1:]:
            value = row.cells[0].text.strip()
            if value.isdigit():
                numbers.append(int(value))
    if numbers != list(range(1, feature_count + 1)):
        raise ValueError("Feature numbering is not continuous")

    with ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        footer_xml = "".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("word/footer") and name.endswith(".xml")
        )
    if document_xml.count("w:tblHeader") < len(SECTIONS) + 1:
        raise ValueError("Repeating table headers are incomplete")
    if " PAGE " not in footer_xml:
        raise ValueError("Page number field is missing")


def main() -> None:
    feature_count = validate_source_data()
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_overview(doc)
    rendered_count = add_feature_sections(doc)
    add_closing(doc, rendered_count)
    if rendered_count != feature_count:
        raise RuntimeError("Feature numbering does not match source data")
    doc.save(OUTPUT)
    validate_output(OUTPUT, feature_count)
    print(f"SAVED={OUTPUT}")
    print(f"FEATURES={feature_count}")
    print(f"SECTIONS={len(SECTIONS)}")


if __name__ == "__main__":
    main()
