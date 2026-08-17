# -*- coding: utf-8 -*-
"""生成《功能清单与报价方案》Word 文档"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GREEN = RGBColor(0x1d, 0x6b, 0x44)
DARK = RGBColor(0x23, 0x29, 0x1f)
GRAY = RGBColor(0x6e, 0x73, 0x68)
HEADER_BG = "1d6b44"
ROWALT_BG = "f1f6ef"

doc = Document()

# ---- 默认字体（含中文）----
def set_font(style_name, size, bold=False, color=None, font="微软雅黑"):
    st = doc.styles[style_name]
    st.font.name = font
    st.font.size = Pt(size)
    st.font.bold = bold
    if color is not None:
        st.font.color.rgb = color
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn('w:eastAsia'), font)

set_font('Normal', 10.5, color=DARK)

# 页边距
for s in doc.sections:
    s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.2)
    s.left_margin = Cm(2.4); s.right_margin = Cm(2.4)

def cn(run, font="微软雅黑"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font)

def para(text="", size=10.5, bold=False, color=None, align=None, after=6, before=0, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    if color is not None:
        r.font.color.rgb = color
    cn(r)
    return p

def heading(text, level=1):
    sizes = {1: 15, 2: 12.5}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.size = Pt(sizes[level]); r.font.bold = True
    r.font.color.rgb = GREEN if level == 1 else DARK
    cn(r)
    # 底部细线（仅一级）
    if level == 1:
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '4'); bottom.set(qn('w:color'), '1d6b44')
        pbdr.append(bottom); pPr.append(pbdr)
    return p

def shade(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color_hex)
    tcPr.append(sh)

def set_cell(cell, text, bold=False, color=None, size=10, align='left', white=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
    p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                   'right': WD_ALIGN_PARAGRAPH.RIGHT}[align]
    r = p.add_run(text)
    r.font.size = Pt(size); r.font.bold = bold
    if white:
        r.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
    elif color is not None:
        r.font.color.rgb = color
    cn(r)

def make_table(headers, rows, widths=None, money_cols=None, total_row=False):
    money_cols = money_cols or []
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        set_cell(hdr[i], h, bold=True, white=True, size=10, align='center')
        shade(hdr[i], HEADER_BG)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        is_total = total_row and ri == len(rows) - 1
        for ci, val in enumerate(row):
            al = 'right' if ci in money_cols else ('center' if ci == 0 and headers[0] in ('序号','') else 'left')
            if ci in money_cols:
                al = 'right'
            set_cell(cells[ci], str(val), bold=is_total, size=10, align=al,
                     color=(GREEN if is_total and ci in money_cols else None))
            if is_total:
                shade(cells[ci], "e7f1ea")
            elif ri % 2 == 1:
                shade(cells[ci], ROWALT_BG)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t

# ======================= 封面 =======================
para("演示原型 · 商务报价", size=10.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=2, before=30)
ttl = doc.add_paragraph(); ttl.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = ttl.add_run("湖南农家乐数字供应链与引流服务平台"); r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = GREEN; cn(r)
ttl.paragraph_format.space_after = Pt(2)
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("功能清单与报价方案"); r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = DARK; cn(r)
sub.paragraph_format.space_after = Pt(18)

para("大中台 ＋ 小前台 ＋ 联盟", size=11.5, bold=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para("甄选好物供应链中台  ＋  N 个农家乐小程序  ＋  农家乐联盟推客平台",
     size=10.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=24)

info = doc.add_table(rows=4, cols=2); info.alignment = WD_TABLE_ALIGNMENT.CENTER
info.style = 'Table Grid'
rows_info = [("牵头单位", "湖南省电子商务协会"),
             ("项目名称", "湖南农家乐数字供应链与引流服务平台"),
             ("方案日期", "2026 年 6 月 24 日"),
             ("报价有效期", "自方案出具之日起 30 日内")]
for i, (k, v) in enumerate(rows_info):
    set_cell(info.rows[i].cells[0], k, bold=True, size=10.5, align='center'); shade(info.rows[i].cells[0], ROWALT_BG)
    set_cell(info.rows[i].cells[1], v, size=10.5, align='left')
    info.rows[i].cells[0].width = Cm(4); info.rows[i].cells[1].width = Cm(11)

doc.add_page_break()

# ======================= 一、项目概述 =======================
heading("一、项目概述")
para("本平台围绕农家乐经营主体，构建集供应链集采、线上经营、联盟引流、直播带货、履约售后与数据服务于一体的"
     "数字化平台，采用「大中台 ＋ 小前台 ＋ 联盟」模式建设，为农家乐引流、为农家乐降本、为农产品拓销、为乡村消费赋能。",
     after=6)
para("平台由三大端组成：", bold=True, after=4)
for t in [
    "①  甄选好物供应链中台（管理后台）——平台核心基础设施，统一管理供应商、商品库、价格体系、订单履约、售后结算、门店与佣金。",
    "②  农家乐独立小程序（移动端）——每家农家乐独立部署的线上门店，承接门店经营、商城、会员、店长选品上架与进货补货。",
    "③  农家乐联盟推客平台（移动端）——面向消费者、推客、主播的流量入口，基于位置推荐、锁粉归因、直播带货与佣金追踪。",
]:
    para(t, after=4)

# ======================= 二、功能清单 =======================
heading("二、功能清单")

modules = [
    ("（一）甄选好物供应链中台", [
        ("数据看板与经营分析", "平台 GMV / 订单 / 入驻商户 / 供应商实时 KPI；近 7 日交易与订单趋势；品类销售占比；热销商品 TOP 榜；待办事项中心；报表导出。"),
        ("供应商管理与资质审核", "供应商邀请入驻、资质核验、审核 / 驳回；供销社渠道资源接入；合作状态与品类筛选管理。"),
        ("商品库与统一选品", "统一商品库与 SKU 管理；中台甄选与农家乐自有商品审核；集采价 / 建议零售价维护；上下架；覆盖农产品、预制菜、食材调料、酒水饮料、文旅伴手礼、民宿用品等多品类。"),
        ("多维价格体系", "集采价、阶梯价（按采购量分档）、区域价、会员价等多维策略配置；阶梯价示例与生效管理，帮助农家乐降低采购成本。"),
        ("订单履约与物流配送", "订单统一接单与供应商自动分配；仓配发货、运单录入、物流轨迹；六段履约流程可视；批量发货；支持小程序 / 直播 / 进货等多来源订单。"),
        ("售后服务与结算分账", "售后工单（破损补寄 / 退货退款 / 质量理赔）；T+1 自动结算分账（含供应商货款）；售后率统计与批量结算。"),
        ("农家乐门店统一管理", "N 个农家乐门店与小程序统一管理；门店入驻、选品上架、经营数据；门店 GMV 与运营状态。"),
    ]),
    ("（二）农家乐独立小程序", [
        ("门店首页与展示", "门店形象、招牌土菜、特色包厢、特色服务、营销公告、一键电话联系。"),
        ("预约预订", "包厢预订、套餐预订；日期 / 场次 / 人数选择；到店提醒；预订记录管理。"),
        ("特产商城与购物车下单", "商品分类筛选；购物车支持改数量 / 删除；储值支付；订单生成与状态跟踪；中台供 / 本店商品标识。"),
        ("会员储值与积分", "会员卡与等级；储值充值（余额实时增减）、积分体系、会员专属价。"),
        ("身份权限与店长工作台", "顾客 / 店员 / 店长三级身份切换；按权限显示经营专区菜单（顾客不可见经营功能）。"),
        ("中台选品加价上架（店长）", "店长从供应链中台甄选商品，自定义本店零售价加价上架，毛利与毛利率实时测算，上架后顾客即可在商城购买，价差即门店利润。"),
        ("进货端订单全流程", "进货车（改量 / 删除）；进货单提交；订单生命周期（已提交→中台接单→仓配发货→配送中→待收货→已完成）；进度时间线；确认收货；再次进货；进货数据统计。"),
        ("分销推广", "分享门店 / 商品赚佣金、推广海报、私域裂变。"),
    ]),
    ("（三）农家乐联盟推客平台", [
        ("附近农家乐与 LBS 推荐", "定位与城市切换；地图与门店分布；按距离 / 好评 / 人气 / 直播 / 可预订筛选；附近门店列表与一键进店（自动携带推客分享参数）。"),
        ("榜单与乡村旅游线路", "人气榜 / 好评榜 / 直播榜；乡村旅游线路推荐与报名。"),
        ("直播带货与预约提醒", "正在直播列表与观看人数；直播预告与开播预约；直播成交由中台统一承接履约。"),
        ("推客中心与钱包提现", "推客等级、佣金钱包、提现（多到账方式）；推广素材、我的粉丝、佣金明细、推客排行。"),
        ("锁粉归因与佣金分享", "分享链接锁粉、订单归因、佣金追踪；热推好物一键分享赚佣金。"),
    ]),
]

idx = 1
for title, feats in modules:
    heading(title, level=2)
    rows = []
    for name, desc in feats:
        rows.append([idx, name, desc]); idx += 1
    make_table(["序号", "功能模块", "功能说明"], rows, widths=[1.4, 4.2, 10.0], money_cols=[])

# ======================= 三、报价明细 =======================
heading("三、报价明细")
para("本项目采用「基础开发费 ＋ 功能模块」方式报价，含一套标准样板店小程序。金额单位：人民币元。", after=8)

heading("3.1  基础开发费", level=2)
para("含：需求梳理与原型确认、整体 UI / UX 设计、前端通用组件与交互框架、账号与多角色权限体系、"
     "数据库与接口基础架构、三端联调、测试与上线部署、项目管理。", after=6)
make_table(["费用项", "说明", "金额（元）"],
           [["平台基础开发费", "项目通用底座 · 一次性", "50,000"]],
           widths=[4.5, 8.5, 3.0], money_cols=[2])

heading("3.2  供应链中台功能模块", level=2)
zt = [["数据看板与经营分析", "8,000"], ["供应商管理与资质审核", "10,000"],
      ["商品库与统一选品", "12,000"], ["多维价格体系", "10,000"],
      ["订单履约与物流配送", "12,000"], ["售后服务与结算分账", "10,000"],
      ["农家乐门店统一管理", "8,000"], ["小计", "70,000"]]
make_table(["功能模块", "金额（元）"], zt, widths=[12.0, 4.0], money_cols=[1], total_row=True)

heading("3.3  农家乐小程序功能模块", level=2)
xc = [["门店首页与展示", "5,000"], ["预约预订（包厢 / 套餐）", "10,000"],
      ["特产商城与购物车下单", "10,000"], ["会员储值与积分", "8,000"],
      ["身份权限与店长工作台", "7,000"], ["中台选品加价上架", "5,000"],
      ["进货端订单全流程", "5,000"], ["小计", "50,000"]]
make_table(["功能模块", "金额（元）"], xc, widths=[12.0, 4.0], money_cols=[1], total_row=True)

heading("3.4  联盟推客平台功能模块", level=2)
lm = [["附近农家乐与 LBS 推荐", "8,000"], ["榜单与乡村旅游线路", "5,000"],
      ["直播带货与预约提醒", "7,000"], ["推客中心与钱包提现", "6,000"],
      ["锁粉归因与佣金分享", "4,000"], ["小计", "30,000"]]
make_table(["功能模块", "金额（元）"], lm, widths=[12.0, 4.0], money_cols=[1], total_row=True)

# ======================= 四、项目总报价 =======================
heading("四、项目总报价")
total = [["基础开发费", "50,000"],
         ["供应链中台功能模块", "70,000"],
         ["农家乐小程序功能模块", "50,000"],
         ["联盟推客平台功能模块", "30,000"],
         ["项目总金额（一次性）", "200,000"]]
make_table(["费用项", "金额（元）"], total, widths=[12.0, 4.0], money_cols=[1], total_row=True)
para("注：上述总金额 20 万元为平台一次性建设费用，含供应链中台、一套标准样板店农家乐小程序、联盟推客平台的开发与上线。",
     size=9.5, color=GRAY, after=6)

# ======================= 五、其他费用与长期服务 =======================
heading("五、其他费用与长期服务")
make_table(["项目", "计费方式", "金额 / 费率"],
           [["农家乐独立小程序部署", "每新增一家独立部署上线（含上架配置与基础培训）", "15,000 元 / 家"],
            ["平台技术服务费", "按平台整体交易流水抽取", "千分之五（5‰）"]],
           widths=[4.2, 8.8, 3.0], money_cols=[2])
para("说明：", bold=True, after=2, before=4)
for t in [
    "1. 总金额 20 万元含一套标准样板店小程序；此后每新增一家农家乐独立部署小程序，按 1.5 万元 / 家收取（可批量优惠，详见合同）。",
    "2. 平台技术服务费按平台整体交易流水的千分之五（5‰）计提，随平台流水按结算周期收取，用于平台运维、迭代与技术保障。",
]:
    para(t, after=3)

# ======================= 六、付款方式与项目周期 =======================
heading("六、付款方式与项目周期（建议，可协商）")
make_table(["阶段", "节点", "比例", "金额（元）"],
           [["首付款", "合同签订", "30%", "60,000"],
            ["开发款", "开发完成 / 验收前", "50%", "100,000"],
            ["尾款", "验收上线", "20%", "40,000"]],
           widths=[3.2, 6.3, 2.5, 4.0], money_cols=[3])
para("项目周期：预计 8 ～ 10 周完成基础版本（含供应链中台、样板店小程序、联盟推客平台）开发与上线，具体以双方排期为准。",
     after=6, before=4)

# ======================= 七、报价说明 =======================
heading("七、报价说明")
notes = [
    "本报价为软件开发与上线服务费用，不含服务器 / 域名 / SSL 证书、短信、直播推流（CDN）、地图 LBS API、"
    "微信小程序认证与支付报备等第三方平台费用，该类费用按实际产生据实结算或由甲方自行承担提供。",
    "交付物含三端系统及源代码部署、操作说明，并提供上线后 1 年免费质保（缺陷修复）；质保范围外的新增功能与迭代另行评估报价。",
    "每新增一家农家乐独立部署小程序收取 1.5 万元 / 家；平台技术服务费按整体交易流水千分之五计提。",
    "本方案为基于演示原型的预估报价，最终功能范围与金额以双方签订的正式合同及需求确认书为准。",
    "本报价有效期为方案出具之日起 30 日内。",
]
for i, n in enumerate(notes, 1):
    para(f"{i}. {n}", after=4)

para("", after=18)
para("（以下空白，签署见正式合同）", size=9.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)

out = "农家乐数字供应链平台-功能清单与报价方案.docx"
doc.save(out)
print("SAVED:", out)
